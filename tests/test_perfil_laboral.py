"""Pruebas deterministas sin Gemini, Supabase ni descarga de modelos."""
from pathlib import Path
import sys
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from servicios.perfil_laboral import (  # noqa: E402
    limpiar_texto_laboral, texto_laboral_candidato, texto_laboral_vacante,
)


def test_datos_personales_no_entran_a_texto_laboral():
    candidato = SimpleNamespace(
        nombre="Ana Rodríguez", edad="23", sexo="Femenino",
        email="ana@example.com", telefono="5512345678",
        perfil_destacado="Analista de datos",
        experiencia="Nombre: Ana Rodríguez\nEdad: 23\nExperiencia en Python y SQL\nCorreo: ana@example.com",
        skills="SQL, Python, Power BI",
        texto_completo="Domicilio, religión, familia, nombre, edad y mucha información privada",
    )
    texto = texto_laboral_candidato(candidato)
    assert "Python" in texto and "SQL" in texto
    for secreto in ("Rodríguez", "23", "Femenino", "5512345678", "ana@example.com", "religión"):
        assert secreto not in texto


def test_rechaza_placeholder_antiguo():
    candidato = SimpleNamespace(
        perfil_destacado="Perfil Operativo",
        experiencia="Trayectoria no especificada en formato legible.",
        skills="Compromiso, Responsabilidad, Adaptabilidad",
    )
    assert texto_laboral_candidato(candidato) == ""


def test_requisitos_demograficos_explicitos_se_eliminan():
    texto = texto_laboral_vacante(
        "Analista contable",
        "Excel avanzado y conciliación bancaria\nSexo: femenino\nEdad: 25 a 30 años\nEstado civil: soltera",
    )
    assert "Excel" in texto
    assert "conciliación" in texto
    assert "femenino" not in texto
    assert "25" not in texto
    assert "soltera" not in texto


def test_limpieza_texto_laboral_conserva_habilidades():
    texto = limpiar_texto_laboral("Correo: test@example.com\nManejo de Excel y Power BI\nTeléfono: 5512345678")
    assert "Excel" in texto
    assert "example" not in texto
    assert "5512345678" not in texto


def test_extractor_reconoce_habilidades_con_acentos_y_csharp():
    from parser import extraer_skills
    skills = extraer_skills("Experiencia en facturación, atención a clientes, C# y Python")
    assert "Facturacion" in skills
    assert "Atencion A Clientes" in skills
    assert "C#" in skills
    assert "Python" in skills


def test_pdf_y_docx_con_tablas_se_extraen_sin_modelo():
    from io import BytesIO
    import fitz
    from docx import Document
    from servicios.extractores_cv import extraer_texto_multiformato

    pdf = fitz.open()
    pagina = pdf.new_page()
    pagina.insert_text((50, 75), "Experiencia laboral en Python y SQL con analisis de datos")
    texto_pdf = extraer_texto_multiformato("prueba.pdf", pdf.tobytes())
    assert "Python" in texto_pdf
    assert "SQL" in texto_pdf
    pdf.close()

    doc = Document()
    tabla = doc.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text = "Habilidades"
    tabla.cell(0, 1).text = "Facturacion y contabilidad empresarial"
    output = BytesIO()
    doc.save(output)
    texto_docx = extraer_texto_multiformato("prueba.docx", output.getvalue())
    assert "contabilidad" in texto_docx


def test_pdf_escaneado_activa_ocr_sin_tesseract_real(monkeypatch):
    from io import BytesIO
    from PIL import Image
    import fitz
    import pytesseract
    from servicios.extractores_cv import extraer_texto_multiformato

    imagen = Image.new("RGB", (250, 90), "white")
    archivo = BytesIO()
    imagen.save(archivo, format="PNG")
    pdf = fitz.open()
    pagina = pdf.new_page()
    pagina.insert_image(fitz.Rect(0, 0, 250, 90), stream=archivo.getvalue())
    datos = pdf.tobytes()
    pdf.close()

    from servicios import extractores_cv

    monkeypatch.setattr(
        extractores_cv,
        "_texto_por_ocr",
        lambda imagen: "Experiencia laboral en limpieza y mantenimiento de oficinas"
    )