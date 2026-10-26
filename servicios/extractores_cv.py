"""Extraccion de texto de PDF, DOCX, TXT e imagenes.

El OCR de imagenes requiere el ejecutable de Tesseract, ademas de pytesseract.
En Windows se busca el ejecutable en PATH y ubicaciones comunes, sin modificar
variables de entorno ni acceder a la base de datos.
"""

import io
import os
import re
import shutil
from pathlib import Path


def limpiar_texto_para_db(texto: str) -> str:
    """Retira bytes nulos y secuencias inutiles antes de guardar texto."""
    if not texto:
        return ""
    t = texto.replace("\x00", " ")
    t = re.sub(r"<<.*?>>", "", t, flags=re.DOTALL)
    t = re.sub(r"/[A-Za-z0-9_]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def localizar_tesseract() -> str:
    """Devuelve la ruta del ejecutable o un error explicito y accionable.

    TESSERACT_CMD permite definir una ubicacion personalizada en el proceso de
    FastAPI. No se escribe en el .env ni se altera el PATH del usuario.
    """
    configurado = os.getenv("TESSERACT_CMD", "").strip().strip('"')
    if configurado:
        ruta = Path(configurado).expanduser()
        if ruta.is_file():
            return str(ruta)
        encontrado = shutil.which(configurado)
        if encontrado:
            return encontrado
        raise ValueError(
            "OCR no disponible: TESSERACT_CMD apunta a un ejecutable inexistente. "
            "Corrige la variable y reinicia FastAPI."
        )

    encontrado = shutil.which("tesseract")
    if encontrado:
        return encontrado

    # Buscar rutas habituales de Tesseract instalado mediante winget / UB Mannheim.
    carpetas = [os.environ.get("ProgramFiles", ""),
                os.environ.get("ProgramFiles(x86)", ""),
                os.environ.get("LOCALAPPDATA", "")]
    carpetas += [r"C:\Program Files", r"C:\Program Files (x86)"]
    for carpeta in carpetas:
        if not carpeta:
            continue
        base = Path(carpeta)
        for candidato in (
            base / "Tesseract-OCR" / "tesseract.exe",
            base / "Programs" / "Tesseract-OCR" / "tesseract.exe",
        ):
            if candidato.is_file():
                return str(candidato)

    raise ValueError(
        "No se encontro Tesseract OCR en el proceso de FastAPI. "
        "Comprueba C:\\Program Files\\Tesseract-OCR\\tesseract.exe "
        "o configura TESSERACT_CMD y reinicia FastAPI."
    )


def _texto_por_ocr(imagen) -> str:
    """Ejecuta OCR sin depender de que el PATH de Streamlit y API coincidan."""
    import pytesseract
    from PIL import ImageOps

    pytesseract.pytesseract.tesseract_cmd = localizar_tesseract()
    if imagen.mode != "RGB":
        imagen = imagen.convert("RGB")
    imagen = ImageOps.autocontrast(imagen)
    try:
        return pytesseract.image_to_string(imagen, config="--psm 3")
    except pytesseract.TesseractNotFoundError as exc:
        raise ValueError(
            "El backend no puede ejecutar Tesseract OCR. "
            "Configura TESSERACT_CMD y reinicia FastAPI."
        ) from exc
    except pytesseract.TesseractError as exc:
        raise ValueError(
            "Tesseract no pudo leer la imagen. Comprueba que el ejecutable "
            "y el idioma de OCR esten instalados."
        ) from exc


def extraer_texto_multiformato(filename: str, file_bytes: bytes) -> str:
    """Extrae texto sin inventar perfiles cuando la lectura falla."""
    ext = filename.rsplit(".", 1)[-1].lower()
    texto = ""
    if ext == "pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            texto = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            pass
        if not texto.strip():
            try:
                import fitz
                with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                    texto = "\n".join(page.get_text() for page in doc)
            except Exception:
                pass
        if len(texto.strip()) < 35:
            try:
                import fitz
                from PIL import Image
                with fitz.open(stream=file_bytes, filetype="pdf") as documento:
                    partes = []
                    for numero in range(min(len(documento), 3)):
                        pix = documento[numero].get_pixmap(
                            matrix=fitz.Matrix(1.5, 1.5), alpha=False
                        )
                        imagen = Image.open(io.BytesIO(pix.tobytes("png")))
                        partes.append(_texto_por_ocr(imagen))
                    texto = "\n".join(partes)
            except ValueError:
                raise
            except Exception as exc:
                raise ValueError(
                    "No se pudo extraer texto del PDF escaneado mediante OCR."
                ) from exc
    elif ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        partes = [p.text for p in doc.paragraphs if p.text.strip()]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(
                    celda.text for celda in fila.cells if celda.text.strip()
                ))
        texto = "\n".join(partes)
    elif ext == "doc":
        raise ValueError(
            "El formato .doc antiguo no esta soportado. Guarda el CV como .docx o PDF."
        )
    elif ext == "txt":
        texto = file_bytes.decode("utf-8-sig", errors="replace")
    elif ext in ("png", "jpg", "jpeg", "webp"):
        try:
            from PIL import Image
            imagen = Image.open(io.BytesIO(file_bytes))
            texto = _texto_por_ocr(imagen)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(
                "No se pudo procesar la imagen. Comprueba que sea valida y "
                "que Tesseract OCR este disponible en FastAPI."
            ) from exc
    else:
        raise ValueError("Formato de CV no admitido")

    return limpiar_texto_para_db(texto)[:15000]
